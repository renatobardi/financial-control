"""Leitura do arquivo de carga: CSV (UTF-8, separador `;` ou `,`) ou planilha `.xlsx` (primeira
aba), sempre com o cabeçalho exato. Falha antes de gravar."""

import csv
import io
import zipfile
from datetime import date, datetime


class ArquivoInvalido(Exception):
    """O arquivo inteiro é recusado. `regra` é o código mostrado no relatório."""

    def __init__(self, regra: str, mensagem: str):
        super().__init__(mensagem)
        self.regra = regra


_SEPARADORES = (";", ",")
_ASSINATURA_XLSX = b"PK\x03\x04"
# O arquivo enviado já é limitado em bytes comprimidos; este é o teto do que ele pode virar
# aberto, contra um zip que expande demais (zip bomb) antes de a planilha ser lida.
LIMITE_DESCOMPRIMIDO = 50 * 1024 * 1024
_LIMITE_LINHAS = 50_000


def _decodificar(conteudo: bytes) -> str:
    try:
        return conteudo.decode("utf-8-sig")
    except UnicodeDecodeError as erro:
        raise ArquivoInvalido(
            "CSV_CODIFICACAO",
            "O arquivo não está em UTF-8. Salve o CSV como UTF-8 e envie de novo.",
        ) from erro


def _cabecalho_ok(celulas: list[str], colunas: tuple[str, ...], opcionais: tuple[str, ...]) -> bool:
    nomes = [c.strip() for c in celulas]
    permitidos = set(colunas) | set(opcionais)
    return set(colunas) <= set(nomes) <= permitidos and len(set(nomes)) == len(nomes)


def _celula_para_texto(valor: object) -> str:
    """O que o CSV teria: número sem `.0` sobrando, data do Excel como `AAAA-MM`."""
    if valor is None:
        return ""
    if isinstance(valor, bool):
        return str(valor)
    if isinstance(valor, int):
        return str(valor)
    if isinstance(valor, float):
        if valor == int(valor):
            return str(int(valor))
        centavos = round(valor * 100)
        return f"{centavos / 100:.2f}" if abs(valor * 100 - centavos) < 1e-6 else repr(valor)
    if isinstance(valor, datetime | date):
        return f"{valor.year:04d}-{valor.month:02d}"
    return str(valor).strip()


def _ler_xlsx(
    conteudo: bytes, colunas: tuple[str, ...], opcionais: tuple[str, ...]
) -> list[tuple[int, dict[str, str]]]:
    from openpyxl import load_workbook  # só quando chega uma planilha

    try:
        with zipfile.ZipFile(io.BytesIO(conteudo)) as zipado:
            if sum(i.file_size for i in zipado.infolist()) > LIMITE_DESCOMPRIMIDO:
                raise ArquivoInvalido("XLSX_GRANDE", "A planilha é grande demais depois de aberta.")
        livro = load_workbook(io.BytesIO(conteudo), read_only=True, data_only=True)
        linhas_brutas = list(livro.worksheets[0].iter_rows(values_only=True))
    except ArquivoInvalido:
        raise
    except Exception as erro:  # zip quebrado, XML inválido, planilha sem abas
        raise ArquivoInvalido(
            "XLSX_INVALIDO", "Não consegui abrir a planilha. Salve como .xlsx e envie de novo."
        ) from erro
    if len(linhas_brutas) > _LIMITE_LINHAS:
        raise ArquivoInvalido("XLSX_GRANDE", f"A planilha passa de {_LIMITE_LINHAS} linhas.")
    cabecalho = [_celula_para_texto(c) for c in (linhas_brutas[0] if linhas_brutas else ())]
    while cabecalho and not cabecalho[-1]:  # colunas em branco à direita
        cabecalho.pop()
    if not cabecalho:
        raise ArquivoInvalido("XLSX_VAZIO", "A planilha está vazia.")
    if not _cabecalho_ok(cabecalho, colunas, opcionais):
        faltando = [c for c in colunas if c not in cabecalho]
        sobrando = [c for c in cabecalho if c and c not in colunas and c not in opcionais]
        partes = []
        if faltando:
            partes.append("faltam as colunas: " + ", ".join(faltando))
        if sobrando:
            partes.append("colunas desconhecidas: " + ", ".join(sobrando))
        detalhe = "; ".join(partes) or "colunas repetidas"
        raise ArquivoInvalido(
            "XLSX_CABECALHO",
            f"Cabeçalho inválido ({detalhe}). Esperado: {', '.join(colunas)}.",
        )
    linhas = []
    for numero, bruta in enumerate(linhas_brutas[1:], start=2):
        textos = [_celula_para_texto(c) for c in bruta[: len(cabecalho)]]
        textos += [""] * (len(cabecalho) - len(textos))
        if not any(textos):
            continue
        linha = dict(zip(cabecalho, textos, strict=True))
        for opcional in opcionais:
            linha.setdefault(opcional, "")
        linhas.append((numero, linha))
    if not linhas:
        raise ArquivoInvalido("XLSX_VAZIO", "A planilha só tem o cabeçalho, sem linhas de dado.")
    return linhas


def ler(
    conteudo: bytes, colunas: tuple[str, ...], opcionais: tuple[str, ...] = ()
) -> list[tuple[int, dict[str, str]]]:
    """Devolve `(número da linha no arquivo, {coluna: valor})` para cada linha com dado."""
    if conteudo[:4] == _ASSINATURA_XLSX:
        return _ler_xlsx(conteudo, colunas, opcionais)
    texto = _decodificar(conteudo)
    if not texto.strip():
        raise ArquivoInvalido("CSV_VAZIO", "O arquivo está vazio.")
    primeira = texto.splitlines()[0]
    escolhido = next(
        (s for s in _SEPARADORES if _cabecalho_ok(primeira.split(s), colunas, opcionais)), None
    )
    if escolhido is None:
        celulas = {c.strip() for sep in _SEPARADORES for c in primeira.split(sep)}
        faltando = [c for c in colunas if c not in celulas]
        sobrando = sorted(
            c for c in celulas if c and c not in colunas and c not in opcionais and len(c) < 40
        )
        partes = []
        if faltando:
            partes.append("faltam as colunas: " + ", ".join(faltando))
        if sobrando:
            partes.append("colunas desconhecidas: " + ", ".join(sobrando))
        detalhe = "; ".join(partes) or "separador inesperado (use ; ou ,)"
        esperado = ", ".join(colunas)
        raise ArquivoInvalido(
            "CSV_CABECALHO",
            f"Cabeçalho inválido ({detalhe}). Esperado, separado por ; ou ,: {esperado}.",
        )
    leitor = csv.reader(io.StringIO(texto), delimiter=escolhido)
    cabecalho = [c.strip() for c in next(leitor)]
    linhas = []
    for numero, celulas in enumerate(leitor, start=2):
        if not any(c.strip() for c in celulas):
            continue
        if len(celulas) != len(cabecalho):
            raise ArquivoInvalido(
                "CSV_LINHA",
                f"A linha {numero} tem {len(celulas)} colunas; o cabeçalho tem {len(cabecalho)}.",
            )
        linha = {c: v.strip() for c, v in zip(cabecalho, celulas, strict=True)}
        for opcional in opcionais:
            linha.setdefault(opcional, "")
        linhas.append((numero, linha))
    if not linhas:
        raise ArquivoInvalido("CSV_VAZIO", "O arquivo só tem o cabeçalho, sem linhas de dado.")
    return linhas
